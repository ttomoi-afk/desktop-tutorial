# -*- coding: utf-8 -*-
"""案件が入っている現物の案件管理表を、build_pipeline.py の現在の列構成へ移し替える。

評価項目を足す・列を並べ替えるなど build_pipeline.py の COLS を変えたあとに使う。
案件管理タブを作り直して値を見出し名で移し、ほかのタブにある案件管理への参照
（質問リストの INDEX/MATCH、仲介会社管理シートの COUNTIFS）を新しい列へ付け替える。

  python3 tools/migrate_layout.py 現物.xlsx -o 新.xlsx [--spec spec.json]

作り直すタブ   案件管理・取込・判定基準・選択肢（判定基準のしきい値は現物の値を引き継ぐ）
そのまま残す   質問リスト・仲介会社管理シート（案件管理への参照だけ付け替える）

spec.json（任意）
  {
    "value_map": {"業種区分": {"旧の値": "新しい値"}},      # 全行に効く置き換え
    "set":     {"16": {"★1-1": "×", "エリア": "◯"}},       # 案件No → 見出し → 値
    "prepend": {"16": {"初期検討メモ": "【10/15】…"}},      # 既存の文の先頭に足す
    "questions": [{"案件No": 10, "分類": "…", "質問（このまま読める文）": "…",
                   "何を確かめたいか": "…", "関連項目": "…", "優先度": "重要"}]
  }
"""
import argparse
import json
import os
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_pipeline as B                                    # noqa: E402
from relayout_live import is_formula, norm, repoint          # noqa: E402

SH = "案件管理"
SH_Q, SH_MED = "質問リスト", "仲介会社管理シート"


def head_map(ws, row):
    out = {}
    for c in ws[row]:
        if c.value is not None:
            out.setdefault(norm(c.value), c.column_letter)
    return out


def harvest(ws, heads):
    """明細行の手入力の値を、行番号ごと・見出し名つきで拾う。"""
    rows = {}
    for r in range(B.FIRST, B.LAST + 1):
        row = {}
        for h, cl in heads.items():
            v = ws[f"{cl}{r}"].value
            if v in (None, "") or is_formula(v):
                continue
            row[h] = v
        if row:
            rows[r] = row
    return rows


def med_last(ws):
    """仲介会社管理シートで数式が用意されている最後の行（ドロップダウンの範囲に使う）。"""
    last = B.MED_FIRST
    for r in range(B.MED_FIRST, ws.max_row + 1):
        if is_formula(ws.cell(r, len(B.MED_ATTRS) + 1).value):
            last = r
    return last


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--spec")
    a = ap.parse_args()
    spec = json.load(open(a.spec, encoding="utf-8")) if a.spec else {}

    wb = load_workbook(a.src)
    order = list(wb.sheetnames)
    old = wb[SH]
    old_heads = head_map(old, B.HEAD_ROW)
    rows = harvest(old, old_heads)
    print(f"案件管理 {len(rows)}行を回収")

    # しきい値は現物の値を引き継ぐ（除外業種は value_map で置き換える）
    old_params = {r: wb["判定基準"][f"B{r}"].value
                  for r in range(B.PARAM_ROW, B.P_EXCL + 1)}

    # 仲介会社のドロップダウンを、管理シートで用意済みの行まで効かせる
    B.MED_LAST = med_last(wb[SH_MED])

    # ── 案件管理を作り直して値を移す ─────────────────────────────────
    idx = wb.sheetnames.index(SH)
    del wb[SH]
    ws = wb.create_sheet(SH, idx)
    B.build_deals(ws, samples=False)
    new_heads = head_map(ws, B.HEAD_ROW)
    gone = [h for h in old_heads if h not in new_heads]
    if gone:
        raise SystemExit(f"新しい列構成に無い旧見出し: {gone}")
    vmap = spec.get("value_map", {})
    nmapped = 0
    for r, row in rows.items():
        for h, v in row.items():
            nv = vmap.get(h, {}).get(v, v) if isinstance(v, str) else v
            nmapped += nv != v
            ws[f"{new_heads[h]}{r}"].value = nv
    print(f"   値を転記（置き換え {nmapped}セル）")

    # ── 案件Noごとの修正 ──────────────────────────────────────────────
    no_row = {}
    for r in range(B.FIRST, B.LAST + 1):
        v = ws[f"{new_heads['No']}{r}"].value
        try:
            no_row[int(float(v))] = r
        except (TypeError, ValueError):
            pass
    for key, kind in (("set", "set"), ("prepend", "prepend")):
        for no, vals in spec.get(key, {}).items():
            r = no_row[int(no)]
            for h, v in vals.items():
                h = norm(h)
                c = ws[f"{new_heads[h]}{r}"]
                if is_formula(c.value):
                    raise SystemExit(f"数式の列は書き換えない: {h}")
                c.value = v if kind == "set" else v + (c.value or "")
    print(f"   案件ごとの修正 {sum(len(v) for v in spec.get('set', {}).values())}"
          f"＋{sum(len(v) for v in spec.get('prepend', {}).values())}セル")

    # ── ほかのタブの参照を付け替える ──────────────────────────────────
    col_map = {old_heads[h]: new_heads[h] for h in old_heads}
    row_map = {B.FIRST: B.FIRST, B.LAST: B.LAST}
    for name in wb.sheetnames:
        if name == SH:
            continue
        n = repoint(wb[name], col_map, row_map)
        if n:
            print(f"参照の付け替え {name}: {n}セル")

    # ── 取込・判定基準・選択肢を作り直す ──────────────────────────────
    for name, build in (("取込", lambda w: B.build_intake(w, B.L)),
                        ("判定基準", B.build_criteria),
                        ("選択肢", B.build_choices)):
        i = wb.sheetnames.index(name)
        del wb[name]
        build(wb.create_sheet(name, i))
    for r, v in old_params.items():
        if isinstance(v, str):
            v = vmap.get("業種区分", {}).get(v, v)
        wb["判定基準"][f"B{r}"].value = v
    print(f"判定基準のしきい値を引き継ぎ {list(old_params.values())}"
          f" → {[wb['判定基準'][f'B{r}'].value for r in old_params]}")

    # ── 追加の質問 ────────────────────────────────────────────────────
    qs = spec.get("questions", [])
    if qs:
        wq = wb[SH_Q]
        qh = head_map(wq, 3)
        used = [r for r in range(4, wq.max_row + 1)
                if wq[f"{qh['案件No']}{r}"].value not in (None, "")]
        nxt = (max(used) if used else 3) + 1
        for q in qs:
            no = q["案件No"]
            seq = max([int(float(wq[f"{qh['Q#']}{r}"].value or 0)) for r in used
                       if str(wq[f"{qh['案件No']}{r}"].value) in (str(no), f"{no}.0")]
                      or [0]) + 1
            vals = dict(q, **{"Q#": seq, "状態": "未質問"})
            for h, v in vals.items():
                h = norm(h)
                if is_formula(wq[f"{qh[h]}{nxt}"].value):
                    raise SystemExit(f"質問リストの数式の列に書こうとした: {h}")
                wq[f"{qh[h]}{nxt}"].value = v
            used.append(nxt)
            nxt += 1
        print(f"質問を追加 {len(qs)}件")

    wb._sheets = [wb[n] for n in order]
    wb.calculation.fullCalcOnLoad = True
    wb.save(a.out)
    print("saved", a.out)


if __name__ == "__main__":
    main()
