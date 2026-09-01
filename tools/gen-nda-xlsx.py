# NDA管理シート（Excel）を組み立てる。列・状態判定の考え方は nda/ のWebアプリと合わせている。
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from datetime import date, timedelta

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "nda", "NDA管理シート.xlsx")
JP = "Meiryo"                    # 日本語のビジネス標準フォント（無い環境ではゴシックにフォールバック）
LAST_ROW = 151                   # 150件ぶんの記入欄（足りなければ最終行をコピーして増やす）

GREEN = "23825A"
INK = "1F2A24"
MUTED = "6B7A72"
LINE = "E3E9E5"

def F(size=11, bold=False, color=INK, italic=False):
    return Font(name=JP, size=size, bold=bold, color=color, italic=italic)

thin = Side(style="thin", color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ══════════════════════════════════════════════════════════
# NDA一覧
# ══════════════════════════════════════════════════════════
ws = wb.active
ws.title = "NDA一覧"

HEAD = [
    ("企業名", 26), ("締結日", 12), ("有効期限", 12), ("自動更新", 9),
    ("状態", 11), ("残日数", 11), ("種別", 11), ("目的・案件名", 26),
    ("自社担当", 14), ("先方窓口", 16), ("原本の保管先", 26), ("備考", 30),
]
for i, (name, width) in enumerate(HEAD, start=1):
    c = ws.cell(row=1, column=i, value=name)
    c.font = F(10.5, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=GREEN)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = box
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[1].height = 28

NOTES = {
    "A1": "契約の相手方。集計の「取引先（社数）」はこの名前で数えます。",
    "B1": "NDAを締結した日。yyyy/m/d で入力してください。",
    "C1": "契約の満了日。自動更新や期限の定めがない場合は空欄のままで構いません。",
    "D1": "申し出がなければ延長される契約は「○」を選びます（状態は「自動更新」になります）。",
    "E1": "数式で自動判定します（有効／期限間近／期限切れ／自動更新／期限未設定）。上書きしないでください。",
    "F1": "有効期限までの残り日数。数式で自動計算します。",
}
for ref, text in NOTES.items():
    ws[ref].comment = Comment(text, "NDA管理シート")

# 記入例（使い方シートの案内どおり、そのまま消して使える）
today = date.today()
samples = [
    ["株式会社あおば商事", date(today.year - 1, 4, 10), date(today.year + 1, 4, 9), "", "相互",
     "共同キャンペーンの検討", "営業部 山田", "経営企画室 佐藤様",
     "共有ドライブ > 法務 > NDA", "（記入例）この行は消してお使いください"],
    ["みなとテクノロジー株式会社", date(today.year, 1, 22), None, "○", "当社開示",
     "受発注システムの改修", "情報システム部 鈴木", "開発部 井上様",
     "法務キャビネット B-3", "（記入例）自動更新の契約は有効期限を空欄に"],
]
for r, row in enumerate(samples, start=2):
    values = [row[0], row[1], row[2], row[3], None, None, row[4], row[5], row[6], row[7], row[8], row[9]]
    for i, v in enumerate(values, start=1):
        if v is not None:
            ws.cell(row=r, column=i, value=v)

# 全行に共通の書式と数式（状態・残日数）を敷いておく
for r in range(2, LAST_ROW + 1):
    ws.cell(row=r, column=5, value=(
        f'=IF($A{r}="","",'
        f'IF($D{r}="○","自動更新",'
        f'IF($C{r}="","期限未設定",'
        f'IF($C{r}<TODAY(),"期限切れ",'
        f'IF($C{r}-TODAY()<=集計!$B$3,"期限間近","有効")))))'
    ))
    # 自動更新の契約は更新日を過ぎても残日数を出さない（超過表示は誤解を招くため）
    ws.cell(row=r, column=6, value=(
        f'=IF(OR($A{r}="",$C{r}=""),"",'
        f'IF(AND($D{r}="○",$C{r}<TODAY()),"",$C{r}-TODAY()))'
    ))
    for i in range(1, len(HEAD) + 1):
        c = ws.cell(row=r, column=i)
        c.font = F(10.5) if i not in (5, 6) else F(10.5, bold=(i == 5))
        c.border = box
        c.alignment = Alignment(vertical="center", wrap_text=(i in (8, 11, 12)))
    ws.cell(row=r, column=2).number_format = "yyyy/mm/dd"
    ws.cell(row=r, column=3).number_format = "yyyy/mm/dd"
    ws.cell(row=r, column=6).number_format = '0"日";0"日超過"'
    for i in (2, 3, 4, 5, 6, 7):
        ws.cell(row=r, column=i).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[r].height = 20

# 入力規則（プルダウン）
dv_auto = DataValidation(type="list", formula1='"○"', allow_blank=True,
                         prompt="自動更新の契約なら「○」を選びます", promptTitle="自動更新")
dv_type = DataValidation(type="list", formula1='"相互,当社開示,先方開示"', allow_blank=True,
                         prompt="どちらが秘密情報を開示するか", promptTitle="種別")
dv_date = DataValidation(type="date", operator="between",
                         formula1="DATE(1990,1,1)", formula2="DATE(2100,12,31)",
                         allow_blank=True, showErrorMessage=True,
                         error="日付として入力してください（例 2026/4/1）", errorTitle="日付の形式")
ws.add_data_validation(dv_auto); dv_auto.add(f"D2:D{LAST_ROW}")
ws.add_data_validation(dv_type); dv_type.add(f"G2:G{LAST_ROW}")
ws.add_data_validation(dv_date); dv_date.add(f"B2:C{LAST_ROW}")

# 状態による行の色分け
rng = f"A2:L{LAST_ROW}"
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'$E2="期限切れ"'], stopIfTrue=False,
    fill=PatternFill("solid", start_color="FBECEB", end_color="FBECEB"),
    font=Font(name=JP, size=10.5, color="B3261E")))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'$E2="期限間近"'], stopIfTrue=False,
    fill=PatternFill("solid", start_color="FDF4E4", end_color="FDF4E4"),
    font=Font(name=JP, size=10.5, color="A85C06")))
ws.conditional_formatting.add(rng, FormulaRule(
    formula=[f'$E2="自動更新"'], stopIfTrue=False,
    fill=PatternFill("solid", start_color="EEF2FB", end_color="EEF2FB")))

ws.auto_filter.ref = f"A1:L{LAST_ROW}"
ws.freeze_panes = "B2"
ws.sheet_view.zoomScale = 110
ws.print_title_rows = "1:1"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0   # 横は1ページに収め、縦は行数ぶん続ける
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ══════════════════════════════════════════════════════════
# 集計
# ══════════════════════════════════════════════════════════
st = wb.create_sheet("集計")
st.column_dimensions["A"].width = 30
st.column_dimensions["B"].width = 14
st.column_dimensions["C"].width = 46

st["A1"] = "NDA 集計"
st["A1"].font = F(16, bold=True, color=GREEN)
st["A2"] = "※ 数値はすべて「NDA一覧」から自動集計しています。黄色いセルだけ書き換えてください。"
st["A2"].font = F(9.5, color=MUTED)

st["A3"] = "「期限間近」とみなす残日数"
st["A3"].font = F(11, bold=True)
st["B3"] = 90
st["B3"].font = Font(name=JP, size=11, bold=True, color="0000FF")
st["B3"].fill = PatternFill("solid", fgColor="FFFF00")
st["B3"].border = box
st["B3"].alignment = Alignment(horizontal="center")
st["C3"] = "この日数以内に有効期限が来る契約を「期限間近」にします（既定 90日）"
st["C3"].font = F(9.5, color=MUTED)

A = f"NDA一覧!$A$2:$A${LAST_ROW}"
B = f"NDA一覧!$B$2:$B${LAST_ROW}"
E = f"NDA一覧!$E$2:$E${LAST_ROW}"
G = f"NDA一覧!$G$2:$G${LAST_ROW}"

rows = [
    ("件数", None, None),
    ("登録件数", f"=COUNTA({A})", "「NDA一覧」の企業名が入っている行数"),
    ("取引先（社数）", f'=IF(COUNTA({A})=0,0,SUMPRODUCT(({A}<>"")/COUNTIF({A},{A}&"")))', "企業名の重複を除いた数"),
    ("今年の締結件数", f'=SUMPRODUCT((YEAR({B})=YEAR(TODAY()))*({A}<>""))', "締結日が今年の契約"),
    ("状態別", None, None),
    ("有効", f'=COUNTIF({E},"有効")', "期限まで余裕がある契約"),
    ("期限間近", f'=COUNTIF({E},"期限間近")', "更新・再締結の検討が必要"),
    ("期限切れ", f'=COUNTIF({E},"期限切れ")', "継続取引があるなら再締結が必要"),
    ("自動更新", f'=COUNTIF({E},"自動更新")', "解約の申し出がなければ継続"),
    ("期限未設定", f'=COUNTIF({E},"期限未設定")', "有効期限の記入漏れがないか確認"),
    ("種別別", None, None),
    ("相互", f'=COUNTIF({G},"相互")', None),
    ("当社開示", f'=COUNTIF({G},"当社開示")', None),
    ("先方開示", f'=COUNTIF({G},"先方開示")', None),
]

r = 5
for label, formula, note in rows:
    if formula is None:                      # セクション見出し
        st.cell(row=r, column=1, value=label).font = F(10.5, bold=True, color=GREEN)
        r += 1
        continue
    st.cell(row=r, column=1, value=label).font = F(11)
    v = st.cell(row=r, column=2, value=formula)
    v.font = F(12, bold=True)
    v.alignment = Alignment(horizontal="center")
    v.border = box
    if note:
        st.cell(row=r, column=3, value=note).font = F(9.5, color=MUTED)
    r += 1

st.cell(row=r + 1, column=1, value="次にやること").font = F(10.5, bold=True, color=GREEN)
st.cell(row=r + 2, column=1,
        value='「NDA一覧」の「状態」列のフィルタで「期限間近」「期限切れ」を選ぶと、対応が必要な契約だけを一覧できます。')
st.cell(row=r + 2, column=1).font = F(10)
st.sheet_view.showGridLines = False
st.page_setup.orientation = "portrait"
st.page_setup.fitToWidth = 1
st.page_setup.fitToHeight = 0
st.sheet_properties.pageSetUpPr.fitToPage = True

# ══════════════════════════════════════════════════════════
# 使い方
# ══════════════════════════════════════════════════════════
hw = wb.create_sheet("使い方")
hw.column_dimensions["A"].width = 20
hw.column_dimensions["B"].width = 86
hw.sheet_view.showGridLines = False
hw.page_setup.orientation = "landscape"
hw.page_setup.fitToWidth = 1
hw.page_setup.fitToHeight = 0
hw.sheet_properties.pageSetUpPr.fitToPage = True

def line(row, a, b, bold=False, color=INK, size=10.5):
    label = hw.cell(row=row, column=1, value=a)
    label.font = F(size, bold=True, color=GREEN if a else INK)
    label.alignment = Alignment(vertical="top")
    c = hw.cell(row=row, column=2, value=b)
    c.font = F(size, bold=bold, color=color)
    c.alignment = Alignment(vertical="top", wrap_text=True)

hw["A1"] = "NDA管理シート — 使い方"
hw["A1"].font = F(16, bold=True, color=GREEN)
hw["A2"] = "秘密保持契約（NDA）の企業名・締結日・有効期限を1枚で管理します。"
hw["A2"].font = F(10.5, color=MUTED)

guide = [
    ("はじめに", "「NDA一覧」の2行目・3行目は記入例です。行ごと削除してからお使いください。"),
    ("入力する列", "A 企業名 ／ B 締結日 ／ C 有効期限 ／ D 自動更新 ／ G 種別 ／ H 目的・案件名 ／ I 自社担当 ／ J 先方窓口 ／ K 原本の保管先 ／ L 備考"),
    ("触らない列", "E 状態 と F 残日数 は数式です。書き換えると自動判定が止まります（151行目まで数式を入れてあります）。"),
    ("必須", "A 企業名 と B 締結日。この2つが入っていれば集計と状態判定が動きます。"),
    ("日付の入れ方", "2026/4/1 のように入力します。C 有効期限は、自動更新や期限の定めがない契約では空欄で構いません。"),
    ("自動更新", "D列で「○」を選ぶと状態は常に「自動更新」になります（期限切れの警告を出しません）。"),
    ("状態の意味", "有効＝期限まで余裕あり ／ 期限間近＝集計シートB3で決めた日数以内（既定90日・オレンジ） ／ 期限切れ＝赤 ／ 自動更新＝青 ／ 期限未設定＝有効期限が空欄"),
    ("色", "行の色は状態から自動でつきます。塗りつぶしを手で変える必要はありません。"),
    ("並べ替え・絞り込み", "1行目の▼（フィルタ）から。期限が近い順に見るときは C 有効期限 で昇順、直近の締結を見るときは B 締結日 で降順。"),
    ("行を増やすには", "151行目より下に足すときは、151行目をコピーして貼り付けると E・F の数式と書式がそのまま入ります。集計シートの参照範囲も広げてください。"),
    ("印刷", "「NDA一覧」は横向き・幅1ページ・見出し行の繰り返しを設定済みです。空行まで印刷されるので、フィルタで必要な行に絞ってから印刷するときれいに収まります。"),
    ("Webアプリと共通", "同じ項目のWeb版（このリポジトリの nda/）があります。Excelを「CSV UTF-8」で保存すると、そのままWeb版に読み込めます。"),
    ("運用のコツ", "月に一度、集計シートの「期限間近」「期限切れ」の件数を見て、更新・再締結の要否を確認してください。"),
]
r = 4
for a, b in guide:
    line(r, a, b)
    # 全角まじりの本文が折り返しても切れないよう、文字数から行の高さを決める
    hw.row_dimensions[r].height = 16 * max(1, -(-len(b) // 42)) + 8
    r += 1

r += 1
hw.cell(row=r, column=1, value="凡例").font = F(11, bold=True, color=GREEN)
r += 1
legend = [
    ("入力するセル", "白い背景のセル", "FFFFFF"),
    ("設定するセル", "黄色いセル（集計シート B3 の日数）", "FFFF00"),
    ("期限切れの行", "赤い背景", "FBECEB"),
    ("期限間近の行", "オレンジの背景", "FDF4E4"),
    ("自動更新の行", "青い背景", "EEF2FB"),
]
for label, text, color in legend:
    c = hw.cell(row=r, column=1, value=label)
    c.font = F(10.5)
    c.fill = PatternFill("solid", fgColor=color)
    c.border = box
    hw.cell(row=r, column=2, value=text).font = F(10.5, color=MUTED)
    r += 1

wb.save(OUT)
print("saved", OUT)
