# -*- coding: utf-8 -*-
"""議事録から拾った回答を、案件管理表の質問リストタブへ書き込む。

  python3 tools/apply_answers.py live.xlsx -a answers.json -o out.xlsx [--dry-run]

answers.json
  {"案件No": 4, "聞いた日": "2026-10-06",
   "出所": "maXアドバイザリー 則兼様との面談（議事録 2026/10/6）",
   "answers": [
     {"Q": 7, "状態": "解決", "回答": "...", "回答で動く評価": "5",
      "回答日": "2026-10-06"},
     {"Q": 6, "状態": "回答待ち", "回答": "..."}        ← 回答日は省略可
   ],
   "関連項目の修正": {"11": "1-3"}                      ← Q番号 → 正しい値
  }

案件Noと Q# で行を引くので、行番号は書かない。質問文・分類・優先度・関連項目は
触らない（関連項目の修正を明示したときだけ直す）。既に回答が入っている行に
別の回答を書こうとしたら止まる。前の回答を残したまま書き足すときは
`"追記": true` を付ける（「／【10/8追記】…」の形で後ろに足し、回答日は新しい日付に）。
"""
import argparse
import datetime
import json
import os

from openpyxl import load_workbook
from openpyxl.styles import Font

SH_Q = "質問リスト"
Q_HEAD, Q_FIRST, Q_LAST = 3, 4, 303
STATES = ["未質問", "回答待ち", "回答済", "追加確認が必要", "解決"]
FONT = Font(name="Arial", size=10, color="141414")


def head_map(ws, row):
    out = {}
    for c in ws[row]:
        if c.value is not None:
            out.setdefault(str(c.value).replace("\n", "").strip(), c.column_letter)
    return out


def as_date(v):
    if v in (None, ""):
        return None
    y, m, d = (int(x) for x in str(v).split("-"))
    return datetime.datetime(y, m, d)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("-a", "--answers", required=True, action="append",
                    help="議事録1本につき1ファイル。複数指定できる")
    ap.add_argument("-o", "--out")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    if not a.dry_run and not a.out:
        ap.error("--out か --dry-run のどちらかが必要")

    wb = load_workbook(a.src)
    ws = wb[SH_Q]
    H = head_map(ws, Q_HEAD)

    for path in a.answers:
        with open(path, encoding="utf-8") as fh:
            spec = json.load(fh)
        no = spec["案件No"]
        asked = as_date(spec.get("聞いた日"))
        src = spec.get("出所", "")
        # 案件No × Q# → 行
        idx = {}
        for r in range(Q_FIRST, Q_LAST + 1):
            if ws.cell(r, 1).value == no:
                idx[ws[f'{H["Q#"]}{r}'].value] = r
        print(f"案件No.{no}: 質問{len(idx)}問 / 回答{len(spec['answers'])}件  ({path})")

        for fix_q, want in (spec.get("関連項目の修正") or {}).items():
            r = idx[int(fix_q)]
            cell = ws[f'{H["関連項目"]}{r}']
            print(f"   Q{fix_q} 関連項目 {cell.value!r} → {want!r}")
            if not a.dry_run:
                cell.value = want

        for ans in spec["answers"]:
            r = idx[ans["Q"]]
            state = ans["状態"]
            if state not in STATES:
                raise SystemExit(f"状態が選択肢外: {state}")
            cur = ws[f'{H["回答"]}{r}'].value
            body = ans["回答"]
            if src and src not in body:
                body = f"{body}【出所】{src}"
            if ans.get("追記"):
                d = as_date(ans.get("回答日")) or asked
                tag = f"【{d.month}/{d.day}追記】" if d else "【追記】"
                body = f"{cur}／{tag}{body}" if cur else f"{tag}{body}"
            elif cur and cur != body and cur != ans["回答"]:
                raise SystemExit(f"Q{ans['Q']}（行{r}）には既に別の回答がある: {cur[:40]}"
                                 "（書き足すなら \"追記\": true）")
            print(f"   Q{ans['Q']:>2}（行{r}） {state:<8} {body[:46]}…")
            if a.dry_run:
                continue
            for h, v in (("回答", body), ("状態", state),
                         ("聞いた日", as_date(ans.get("聞いた日")) or asked),
                         ("回答日", as_date(ans.get("回答日"))),
                         ("回答で動く評価", ans.get("回答で動く評価"))):
                if v is None:
                    continue
                c = ws[f"{H[h]}{r}"]
                c.value = v
                c.font = FONT
                if h in ("聞いた日", "回答日"):
                    c.number_format = "yyyy/mm/dd"

    if a.dry_run:
        print("（--dry-run なので書き込んでいない）")
        return
    wb.save(a.out)
    open_now = sum(1 for r in range(Q_FIRST, Q_LAST + 1)
                   if ws.cell(r, 1).value is not None
                   and ws[f'{H["状態"]}{r}'].value != "解決")
    print(f"saved {a.out} / 未解決の質問 {open_now}問")


if __name__ == "__main__":
    main()
