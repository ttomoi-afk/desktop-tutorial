# -*- coding: utf-8 -*-
"""ワークブックに書いた数式を実際に評価する最小エンジン。

LibreOffice がこのコンテナで起動しないため recalc.py が使えない。その代替として
使っている関数だけを実装し、セルの値を本当に計算して期待値と突き合わせる。
対応: IF AND OR NOT COUNTIF COUNTIFS SUM DATE INDEX MATCH IFERROR LEFT EXACT
"""
import datetime
import re
from openpyxl.utils import column_index_from_string, get_column_letter

TOKEN = re.compile(r"""
    (?P<str>"(?:[^"]|"")*")
  | (?P<num>\d+(?:\.\d+)?)
  | (?P<ref>(?:(?:'[^']+'|[^\s!(),:+\-*/=<>"&]+)!)?\$?[A-Za-z][A-Za-z0-9_.]*\$?\d*)
  | (?P<op><=|>=|<>|[+\-*/=<>&])
  | (?P<lp>\()
  | (?P<rp>\))
  | (?P<comma>,)
  | (?P<colon>:)
  | (?P<ws>\s+)
""", re.X)


class Err(Exception):
    def __init__(self, kind):
        self.kind = kind
        super().__init__(kind)


def tokenize(f):
    out, i = [], 0
    while i < len(f):
        m = TOKEN.match(f, i)
        if not m:
            raise Err(f"#PARSE! at {i}: {f[i:i+20]!r}")
        i = m.end()
        k = m.lastgroup
        if k != "ws":
            out.append((k, m.group()))
    return out


class Parser:
    """再帰下降。関数名は ref として読まれるので '(' が続くかで判別する。"""

    def __init__(self, toks, wb, sheet):
        self.t, self.i, self.wb, self.sheet = toks, 0, wb, sheet

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def take(self):
        v = self.t[self.i]
        self.i += 1
        return v

    def parse(self):
        v = self.cmp()
        if self.i != len(self.t):
            raise Err(f"#PARSE! trailing {self.t[self.i:]}")
        return v

    def cmp(self):
        left = self.add()
        while self.peek()[0] == "op" and self.peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            o = self.take()[1]
            right = self.add()
            left = compare(left, o, right)
        return left

    def add(self):
        left = self.mul()
        while self.peek()[0] == "op" and self.peek()[1] in ("+", "-", "&"):
            o = self.take()[1]
            right = self.mul()
            if o == "&":
                left = text(left) + text(right)
            else:
                left = num(left) + num(right) if o == "+" else num(left) - num(right)
        return left

    def mul(self):
        left = self.unary()
        while self.peek()[0] == "op" and self.peek()[1] in ("*", "/"):
            o = self.take()[1]
            right = self.unary()
            if o == "*":
                left = num(left) * num(right)
            else:
                d = num(right)
                if d == 0:
                    raise Err("#DIV/0!")
                left = num(left) / d
        return left

    def unary(self):
        if self.peek()[0] == "op" and self.peek()[1] == "-":
            self.take()
            return -num(self.unary())
        return self.atom()

    def atom(self):
        k, v = self.take()
        if k == "num":
            return float(v)
        if k == "str":
            return v[1:-1].replace('""', '"')
        if k == "lp":
            r = self.cmp()
            assert self.take()[0] == "rp"
            return r
        if k == "ref":
            if self.peek()[0] == "lp":
                return self.call(v.upper())
            if v.upper() in ("TRUE", "FALSE"):
                return v.upper() == "TRUE"
            if self.peek()[0] == "colon":
                self.take()
                k2, v2 = self.take()
                return self.wb.range(self.sheet, v, v2)
            return self.wb.cell(self.sheet, v)
        raise Err(f"#PARSE! unexpected {k} {v}")

    def args(self):
        assert self.take()[0] == "lp"
        out = []
        if self.peek()[0] == "rp":
            self.take()
            return out
        while True:
            out.append(("lazy", self.i))
            depth = 0
            while True:
                k, v = self.peek()
                if k is None:
                    raise Err("#PARSE! eof")
                if k == "lp":
                    depth += 1
                elif k == "rp":
                    if depth == 0:
                        break
                    depth -= 1
                elif k == "comma" and depth == 0:
                    break
                self.take()
            out[-1] = (out[-1][1], self.i)
            k, _ = self.take()
            if k == "rp":
                break
        return out

    def ev(self, span):
        sub = Parser(self.t[span[0]:span[1]], self.wb, self.sheet)
        return sub.parse()

    def call(self, name):
        a = self.args()
        if name == "IF":
            c = truth(self.ev(a[0]))
            if c:
                return self.ev(a[1])
            return self.ev(a[2]) if len(a) > 2 else False
        if name == "IFERROR":
            try:
                return self.ev(a[0])
            except Err:
                return self.ev(a[1])
        if name == "AND":
            return all(truth(self.ev(x)) for x in a)
        if name == "OR":
            return any(truth(self.ev(x)) for x in a)
        if name == "NOT":
            return not truth(self.ev(a[0]))
        if name == "SUM":
            t = 0.0
            for x in a:
                for v in flat(self.ev(x)):
                    if v is None or v == "" or isinstance(v, str):
                        continue
                    t += num(v)
            return t
        if name == "DATE":
            return datetime.date(int(num(self.ev(a[0]))), int(num(self.ev(a[1]))),
                                 int(num(self.ev(a[2]))))
        if name == "COUNTIFS":
            if len(a) % 2:
                raise Err("#VALUE! COUNTIFS")
            pairs = [(flat(self.ev(a[i])), self.ev(a[i + 1]))
                     for i in range(0, len(a), 2)]
            n = len(pairs[0][0])
            if any(len(r) != n for r, _ in pairs):
                raise Err("#VALUE! COUNTIFS ranges differ")
            return float(sum(1 for i in range(n)
                             if all(countif_match(r[i], c) for r, c in pairs)))
        if name == "COUNTIF":
            rng, crit = self.ev(a[0]), self.ev(a[1])
            return float(sum(1 for c in flat(rng) if countif_match(c, crit)))
        if name == "EXACT":
            return text(self.ev(a[0])) == text(self.ev(a[1]))
        if name == "LEFT":
            s = text(self.ev(a[0]))
            n = int(num(self.ev(a[1]))) if len(a) > 1 else 1
            return s[:n]
        if name == "MATCH":
            want, rng = self.ev(a[0]), flat(self.ev(a[1]))
            for idx, c in enumerate(rng, start=1):
                if norm(c) == norm(want):
                    return float(idx)
            raise Err("#N/A")
        if name == "INDEX":
            rng = flat(self.ev(a[0]))
            n = int(num(self.ev(a[1])))
            if n < 1 or n > len(rng):
                raise Err("#REF!")
            v = rng[n - 1]
            return "" if v is None else v
        raise Err(f"#NAME? {name}")


EPOCH = datetime.date(1899, 12, 30)      # Excel のシリアル値の起点


def to_serial(v):
    if isinstance(v, datetime.datetime):
        d = v.date()
        frac = (v.hour * 3600 + v.minute * 60 + v.second) / 86400
        return (d - EPOCH).days + frac
    return (v - EPOCH).days


def is_date(v):
    return isinstance(v, (datetime.datetime, datetime.date))


def flat(x):
    return x if isinstance(x, list) else [x]


def norm(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return v
    if isinstance(v, float) and v.is_integer():
        return v
    return v


def num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if is_date(v):
        return float(to_serial(v))
    if isinstance(v, (int, float)):
        return float(v)
    raise Err("#VALUE!")


def text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if is_date(v):
        # Excel の "&" は日付をシリアル値の文字列にする（">="&DATE(...) がこれ）
        return str(int(to_serial(v)))
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def truth(v):
    if isinstance(v, bool):
        return v
    if v is None or v == "":
        return False
    if isinstance(v, (int, float)):
        return v != 0
    raise Err("#VALUE!")


def compare(a, o, b):
    if isinstance(a, str) or isinstance(b, str):
        if isinstance(a, (int, float)) or isinstance(b, (int, float)):
            # Excel: 数値 < 文字列
            an, bn = (0, 1) if isinstance(a, (int, float)) else (1, 0)
            return {"=": an == bn, "<>": an != bn, "<": an < bn,
                    ">": an > bn, "<=": an <= bn, ">=": an >= bn}[o]
        a, b = text(a), text(b)
    else:
        a, b = num(a), num(b)
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b,
            "<=": a <= b, ">=": a >= b}[o]


def countif_match(cell, crit):
    if isinstance(crit, str):
        m = re.match(r"^(<=|>=|<>|<|>|=)(.*)$", crit)
        if m:
            o, rest = m.group(1), m.group(2)
            try:
                rest = float(rest)
            except ValueError:
                pass
            if cell is None and not isinstance(rest, float):
                cell = ""
            try:
                return compare(norm(cell), o, rest)
            except Err:
                return False
        if "*" in crit or "?" in crit:
            raise Err("#WILDCARD-NOT-SUPPORTED")
        return text(norm(cell)) == crit
    return norm(cell) == norm(crit)


class Book:
    """openpyxl の値と、こちらで差し込んだ値を混ぜて読む。数式は再帰評価する。"""

    def __init__(self, wb, maxrow=200):
        self.wb, self.maxrow, self.override = wb, maxrow, {}
        self.depth = 0

    def set(self, sheet, addr, value):
        self.override[(sheet, addr.replace("$", ""))] = value

    def raw(self, sheet, addr):
        a = addr.replace("$", "")
        if (sheet, a) in self.override:
            return self.override[(sheet, a)]
        return self.wb[sheet][a].value

    def cell(self, sheet, ref):
        if "!" in ref:
            sheet, ref = ref.split("!", 1)
            sheet = sheet.strip("'")
        a = ref.replace("$", "")
        if not re.search(r"\d", a):          # 列だけの参照は範囲扱い
            return self.range(sheet, ref, ref)
        v = self.raw(sheet, a)
        if isinstance(v, str) and v.startswith("="):
            self.depth += 1
            if self.depth > 40:
                raise Err("#CIRCULAR!")
            try:
                return evaluate(self, sheet, v)
            finally:
                self.depth -= 1
        return v

    def range(self, sheet, r1, r2):
        if "!" in r1:
            sheet, r1 = r1.split("!", 1)
            sheet = sheet.strip("'")
        if "!" in r2:
            r2 = r2.split("!", 1)[1]
        def split(a):
            a = a.replace("$", "")
            m = re.match(r"^([A-Z]{1,3})(\d*)$", a)
            return m.group(1), (int(m.group(2)) if m.group(2) else None)
        c1, row1 = split(r1)
        c2, row2 = split(r2)
        if row1 is None:
            row1, row2 = 1, self.maxrow
        out = []
        for ci in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
            for ri in range(row1, row2 + 1):
                out.append(self.cell(sheet, f"{get_column_letter(ci)}{ri}"))
        return out


def evaluate(book, sheet, formula):
    f = formula[1:] if formula.startswith("=") else formula
    return Parser(tokenize(f), book, sheet).parse()
