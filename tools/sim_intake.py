# -*- coding: utf-8 -*-
"""gas/案件取込.gs のロジックを Python で再現し、実際のブックで取込を通す。

LibreOffice も Apps Script もこの環境では動かせないため、スクリプトが行う
「見出し名で対応を取り、数式列を避けて書き込み、案件Noを振って質問へ伝播する」
挙動を同じ手順で再現して、結果の数式まで評価して確かめる。
"""
import sys
from openpyxl import load_workbook
from xlformula import Book, evaluate

SH_DEAL, SH_Q, SH_IN = "案件管理", "質問リスト", "取込"
DEAL_HEAD_ROW, DEAL_FIRST = 4, 5
Q_HEAD_ROW, Q_FIRST = 3, 4
IN_DEAL_HEAD, IN_DEAL_ROW = 5, 6
IN_Q_HEAD, IN_Q_FIRST, IN_Q_LAST = 9, 10, 59


def headers(ws, row):
    out = {}
    for c in ws[row]:
        if c.value is None:
            continue
        h = str(c.value).replace("\n", "").strip()
        out.setdefault(h, c.column)
    return out


def is_formula(ws, row, col):
    v = ws.cell(row, col).value
    return isinstance(v, str) and v.startswith("=")


def import_deal(wb, overwrite_answer="new"):
    shIn, shD, shQ = wb[SH_IN], wb[SH_DEAL], wb[SH_Q]
    inH, dH, qH = (headers(shIn, IN_DEAL_HEAD), headers(shD, DEAL_HEAD_ROW),
                   headers(shQ, Q_HEAD_ROW))
    iqH = headers(shIn, IN_Q_HEAD)

    payload = {}
    for h, col in inH.items():
        v = shIn.cell(IN_DEAL_ROW, col).value
        if v not in ("", None):
            payload[h] = v
    if not payload:
        raise SystemExit("■案件 が空")
    name = payload.get("企業名")
    if not name:
        raise SystemExit("企業名が空")

    no_col, name_col = dH["No"], dH["企業名"]
    existing = 0
    for r in range(DEAL_FIRST, shD.max_row + 1):
        if str(shD.cell(r, name_col).value or "").strip() == str(name).strip():
            existing = r
            break
    if existing and overwrite_answer == "overwrite":
        target, deal_no, mode = existing, shD.cell(existing, no_col).value, "上書き"
    else:
        target = DEAL_FIRST
        mx = 0
        for r in range(DEAL_FIRST, shD.max_row + 1):
            if shD.cell(r, no_col).value not in ("", None):
                target = r + 1
                try:
                    mx = max(mx, float(shD.cell(r, no_col).value))
                except (TypeError, ValueError):
                    pass
        deal_no, mode = mx + 1, "新規追加"

    shD.cell(target, no_col).value = deal_no
    wrote = 0
    for h, v in payload.items():
        col = dH.get(h)
        if not col or is_formula(shD, target, col):
            continue
        shD.cell(target, col).value = v
        wrote += 1

    rows = []
    qcol = iqH.get("質問（このまま読める文）", 2)
    for r in range(IN_Q_FIRST, IN_Q_LAST + 1):
        q = shIn.cell(r, qcol).value
        if q in ("", None):
            continue
        rows.append({h: shIn.cell(r, iqH[h]).value for h in iqH})

    added = 0
    if rows:
        qno, qseq = qH["案件No"], qH["Q#"]
        seq, write_row = 0, Q_FIRST
        for r in range(Q_FIRST, shQ.max_row + 1):
            if shQ.cell(r, qno).value not in ("", None):
                write_row = r + 1
                if float(shQ.cell(r, qno).value) == float(deal_no):
                    try:
                        seq = max(seq, float(shQ.cell(r, qseq).value))
                    except (TypeError, ValueError):
                        pass
        for k, row in enumerate(rows):
            rr = write_row + k
            shQ.cell(rr, qno).value = deal_no
            shQ.cell(rr, qseq).value = int(seq) + k + 1
            for h, v in row.items():
                c = qH.get(h)
                if c and v not in ("", None):
                    shQ.cell(rr, c).value = v
            if "状態" in qH:
                shQ.cell(rr, qH["状態"]).value = "未質問"
            added += 1

    for c in range(1, shIn.max_column + 1):
        shIn.cell(IN_DEAL_ROW, c).value = None
        for r in range(IN_Q_FIRST, IN_Q_LAST + 1):
            shIn.cell(r, c).value = None
    return dict(mode=mode, row=target, no=deal_no, wrote=wrote, added=added,
                write_row=write_row if rows else None)
