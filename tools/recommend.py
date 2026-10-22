#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""推薦案件タブ（上司＝服部さんへ推薦する案件を横に並べたまとめ）を作る・案件を足す。

1列＝1案件。「案件No」の行に番号を入れると、案件管理と質問リストから
企業名・判定・13項目・価格（TE定義の①②倍率と目線の株式価値）・次アクションが入る。
手で入れるのは黄色のセルだけ
（推薦日・推薦ポイント・主な論点・推薦理由・服部判断・服部コメント・承継コストの提示額）。
列は30件分を用意してある。足りなくなったら最後の列を右へコピーすれば同じ式が付く。

  # 現物（ドライブからダウンロードした xlsx）にタブを足し、最初の案件を入れる
  python3 tools/recommend.py tab --layout live.xlsx --deals rec.json --out 推薦案件タブ.xlsx --verify

  # タブがある現物に、案件を足す貼り付けブロックを作る（空いている列に入る）
  python3 tools/recommend.py blocks --layout live.xlsx --deals rec.json --out blocks.txt --verify

rec.json
  {"推薦日": "2026/10/01",
   "deals": [{"案件No": 12, "推薦ポイント": "・…\\n・…", "主な論点": "・…",
              "承継コスト": 45.0}]}
  承継コストは仲介の試算などがあるときだけ書く（空欄ならレーマン方式＋DDの概算を使う）。
  推薦理由・服部判断・服部コメントは友井さん・服部さんが書く欄なので、ここからは入れない。
"""
import argparse
import datetime
import io
import json
import os
import re
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_pipeline as B                                    # noqa: E402

SH = "推薦案件"
SH_DEAL, SH_Q = "案件管理", "質問リスト"
DEAL_REF, Q_REF = "'案件管理'", "'質問リスト'"
NCOL = 30                       # 用意する案件の列数（B〜AE）
FIRST_COL = 2
R_LO, R_HI, R_HEAD = 3, 4, 5    # TE目線の倍率（低い方・高い方）と見出し
MONEY, MULT, DATE = "#,##0.0", '0.0"倍"', "yyyy/mm/dd"

# 13項目の短い名前。build_pipeline.ITEMS と食い違ったら落とす（項目を足したらここも足す）
SHORT = {"1-1": "業種ターゲットとの整合", "1-2": "1件あたりの資金負担",
         "1-3": "非キャッシュ性資産", "2": "安定黒字か", "3": "付加価値（粗利率）",
         "4": "永続性", "5": "価格（②倍率）", "8-1": "ストック型かフロー型か",
         "8-2": "流行り廃り", "9": "売却理由", "11": "特定取引先への依存",
         "12": "特定人材への依存", "エリア": "所在エリア（首都圏・関西か）"}
assert list(SHORT) == B.ITEM_NAMES, "build_pipeline.ITEMS と SHORT の項目が合わない"


def norm(h):
    return re.sub(r"\s", "", str(h))


def is_formula(v):
    return isinstance(v, ArrayFormula) or (isinstance(v, str) and v.startswith("="))


def letters(ws, row):
    """見出し行の 見出し（改行を詰めたもの）→ 列記号。"""
    out = {}
    for c in ws[row]:
        if c.value is not None:
            out.setdefault(norm(c.value), c.column_letter)
    return out


# ── 行の並び ─────────────────────────────────────────────────────────────
# (キー, ラベル, 種別, 引数)
#   in_*   手入力（黄色）。no/date/point/issue が Claude の貼り付けブロックの宛先で、連続させる
#   look   案件管理から引く（引数＝案件管理の見出し）。text/num/date で空欄の扱いが違う
#   calc   この表の中で計算
def spec():
    rows = [
        ("no", "案件No", "in_no", None),
        ("date", "推薦日", "in_date", None),
        ("point", "推薦ポイント", "in_text", None),
        ("issue", "主な論点・リスク", "in_text", None),
        ("why", "推薦理由（友井）", "in_text", None),
        ("judge", "服部判断", "in_judge", None),
        ("comment", "服部コメント", "in_text", None),
        ("row", "案件管理の行（自動）", "helper", None),
        ("b1", "基本情報", "band", "7B8B9A"),
        ("name", "企業名", "name", "企業名"),
        ("biz", "事業内容", "text", "事業内容"),
        ("cat", "業種区分", "text", "業種区分"),
        ("loc", "所在地", "text", "所在地"),
        ("med", "仲介会社", "text", "仲介会社"),
        ("medp", "仲介担当者", "text", "仲介担当者"),
        ("inflow", "流入日", "date", "流入日"),
        ("iko", "意向表明期限", "date", "意向表明期限"),
        ("status", "ステータス", "text", "ステータス"),
        ("b2", "判定", "band", "2E6B4F"),
        ("ai", "AI総合判定", "text", "AI総合判定"),
        ("tomoi", "友井判定", "text", "友井判定"),
        ("rate", "達成率（13項目）", "num", ("達成率", "0%")),
    ]
    for name, _, must, *_ in B.ITEMS:
        star = "★" if must else ""
        rows.append((f"g{name}", f"{star}{name}　{SHORT[name]}", "grade", f"{star}{name}"))
    rows += [
        ("b3", "価格（TE定義・百万円）", "band", "3D6B8E"),
        ("px", "希望株式価値（譲渡価格）", "num", ("譲渡価格", MONEY)),
        ("eb", "実態EBITDA", "num", ("実態EBITDA", MONEY)),
        ("sales", "売上", "num", ("売上", MONEY)),
        ("nc", "簿価NetCash", "num", ("簿価NetCash", MONEY)),
        ("wc", "平均必要運転資金", "num", ("平均必要運転資金", MONEY)),
        ("rnc", "実質NetCash（＝簿価NetCash−運転資金）", "calc", MONEY),
        ("ev", "EV（＝株式価値−実質NetCash）", "calc", MONEY),
        ("cost_in", "承継コスト（仲介の試算など。空欄なら概算）", "in_num", None),
        ("cost", "承継コスト（計算に使う額）", "calc", MONEY),
        ("m1", "①EV／EBITDA", "calc", MULT),
        ("m2", "②（EV＋承継コスト）／EBITDA　※No.5の判定", "calc", MULT),
        ("lo", None, "calc", MONEY),
        ("hi", None, "calc", MONEY),
        ("b4", "進行", "band", "6B7F8C"),
        ("act", "次アクション", "text", "次アクション"),
        ("due", "期限", "date", "期限"),
        ("openq", "未解決の質問（件）", "calc", "0"),
        ("mustq", "うち必須（件）", "calc", "0"),
    ]
    return rows


SPEC = spec()
ROW = {k: R_HEAD + 1 + i for i, (k, *_) in enumerate(SPEC)}
LAST_ROW = R_HEAD + len(SPEC)
NOTE_ROW = LAST_ROW + 2
# Claude の貼り付けブロックの宛先（連続していること）
PASTE_KEYS = ["no", "date", "point", "issue"]
assert [ROW[k] for k in PASTE_KEYS] == list(range(ROW["no"], ROW["no"] + 4))


def lehman(p):
    """cash_need.lehman と同じレーマン方式（百万円）。10,000超の部分は数えない。"""
    return (f"IF({p}<=500,{p}*0.05,IF({p}<=1000,25+({p}-500)*0.04,"
            f"IF({p}<=5000,45+({p}-1000)*0.03,IF({p}<=10000,165+({p}-5000)*0.02,265))))")


def dd(p):
    """cash_need.dd_estimate と同じDD費用の概算。"""
    return f"IF({p}<=300,10,IF({p}<=1000,20,IF({p}<=3000,30,45)))"


def formula(key, kind, arg, c, D, Q):
    """列 c（記号）の、キー key の行に入れる式。D・Q は案件管理・質問リストの見出し→列記号。"""
    r = {k: f"{c}${v}" for k, v in ROW.items()}
    rng = lambda h: f"{DEAL_REF}!${D[h]}:${D[h]}"          # noqa: E731
    at = r["row"]
    if kind == "helper":
        return (f'=IF({r["no"]}="","",IFERROR(MATCH({r["no"]},'
                f'{DEAL_REF}!${D["No"]}:${D["No"]},0),""))')
    if kind == "name":
        return (f'=IF({r["no"]}="","",IF({at}="","（案件管理に無い番号）",'
                f'INDEX({rng(arg)},{at})&""))')
    if kind in ("text", "grade"):
        return f'=IF({at}="","",INDEX({rng(arg)},{at})&"")'
    if kind in ("num", "date"):
        h = arg[0] if isinstance(arg, tuple) else arg
        return f'=IF({at}="","",IF(INDEX({rng(h)},{at})="","",INDEX({rng(h)},{at})))'
    px, eb, nc, wc = r["px"], r["eb"], r["nc"], r["wc"]
    rnc, ev, cost = r["rnc"], r["ev"], r["cost"]
    if key == "rnc":
        return f'=IF(OR({nc}="",{wc}=""),"",{nc}-{wc})'
    if key == "ev":
        return f'=IF(OR({px}="",{rnc}=""),"",{px}-{rnc})'
    if key == "cost":
        L = lehman(px)
        return (f'=IF({r["cost_in"]}<>"",{r["cost_in"]},IF({px}="","",'
                f'IF({L}<20,20,{L})+{dd(px)}))')
    if key == "m1":
        return f'=IF(OR({ev}="",{eb}=""),"",IF({eb}<=0,"",{ev}/{eb}))'
    if key == "m2":
        return f'=IF(OR({ev}="",{eb}="",{cost}=""),"",IF({eb}<=0,"",({ev}+{cost})/{eb}))'
    if key in ("lo", "hi"):
        m = f"$A${R_LO if key == 'lo' else R_HI}"
        return f'=IF(OR({eb}="",{rnc}="",{cost}=""),"",{m}*{eb}-{cost}+{rnc})'
    qa = f"{Q_REF}!${Q['案件No']}:${Q['案件No']}"
    qs = f"{Q_REF}!${Q['状態']}:${Q['状態']}"
    if key == "openq":
        return f'=IF({r["no"]}="","",COUNTIFS({qa},{r["no"]},{qs},"<>解決"))'
    if key == "mustq":
        qp = f"{Q_REF}!${Q['優先度']}:${Q['優先度']}"
        return f'=IF({r["no"]}="","",COUNTIFS({qa},{r["no"]},{qp},"必須",{qs},"<>解決"))'
    raise ValueError(key)


def build_recommend(ws, D, Q, ncol=NCOL):
    """推薦案件タブを作る。D・Q は案件管理・質問リストの 見出し→列記号（改行は詰めてよい）。"""
    D = {norm(k): v for k, v in D.items()}
    Q = {norm(k): v for k, v in Q.items()}
    cols = [get_column_letter(FIRST_COL + i) for i in range(ncol)]
    last = cols[-1]

    ws["A1"] = "推薦案件（服部さんへ推薦する案件のまとめ）"
    ws["A1"].font = B.TITLE_F
    ws["A2"] = ("使い方：空いている列の「案件No」に番号を入れると、案件管理と質問リストから自動で入ります。"
                "手で入れるのは黄色のセルだけ。列が足りなくなったら最後の列を右へコピーしてください。")
    ws["A2"].font = B.NOTE_F
    for r, v, label in ((R_LO, 5, "低い方"), (R_HI, 7, "高い方")):
        c = ws[f"A{r}"]
        c.value = v
        c.number_format = f'"TE目線の倍率（{label}）　"0.0"倍"'
        c.font = Font(name=B.FONT, size=10, bold=True, color=B.BLUE_TXT)
        c.fill = B.IN_FILL
        c.border = B.BOX
    ws[f"B{R_LO}"] = "← 倍率を変えると「TE目線の株式価値」が変わります（②＝（EV＋承継コスト）÷実態EBITDA がこの倍率になる株式価値）"
    ws[f"B{R_LO}"].font = B.NOTE_F

    ws[f"A{R_HEAD}"] = "項目"
    for c in ["A"] + cols:
        cell = ws[f"{c}{R_HEAD}"]
        if c != "A":
            cell.value = (f'=IF({c}${ROW["no"]}="","","No."&{c}${ROW["no"]}&'
                          f'IF({c}${ROW["name"]}="","","　"&{c}${ROW["name"]}))')
        cell.font = B.BAND_F
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = B.BOX

    judge_dv = DataValidation(type="list", formula1=f'"{",".join(B.ESC_JUDGE)}"', allow_blank=True)
    ws.add_data_validation(judge_dv)
    judge_dv.add(f"B{ROW['judge']}:{last}{ROW['judge']}")

    for key, label, kind, arg in SPEC:
        r = ROW[key]
        a = ws[f"A{r}"]
        if key in ("lo", "hi"):
            label = f'="TE目線の株式価値（②"&$A${R_LO if key == "lo" else R_HI}&"倍）"'
        a.value = label
        a.border = B.BOX
        a.alignment = Alignment(vertical="top", wrap_text=True)
        if kind == "band":
            a.font = B.BAND_F
            for c in ["A"] + cols:
                ws[f"{c}{r}"].fill = PatternFill("solid", fgColor=arg)
            continue
        a.font = Font(name=B.FONT, size=9, bold=kind.startswith("in") or key in ("name", "ai", "m2"),
                      color=B.GREY if kind == "helper" else B.INK)
        a.fill = B.IN_FILL if kind.startswith("in") else PatternFill("solid", fgColor="F3F3F3")
        for c in cols:
            cell = ws[f"{c}{r}"]
            cell.border = B.BOX
            if kind.startswith("in"):
                cell.fill = B.IN_FILL
                cell.font = Font(name=B.FONT, size=10, color=B.BLUE_TXT if kind == "in_num" else B.INK)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if kind == "in_date":
                    cell.number_format = DATE
                elif kind == "in_num":
                    cell.number_format = MONEY
                continue
            cell.value = formula(key, kind, arg, c, D, Q)
            cell.font = Font(name=B.FONT, size=10 if kind != "helper" else 8,
                             bold=key in ("name", "ai", "tomoi", "m2"),
                             color=B.GREY if kind == "helper" else B.INK)
            cell.alignment = Alignment(vertical="top", wrap_text=kind in ("text", "name"),
                                       horizontal="center" if kind == "grade" else None)
            fmt = arg[1] if kind == "num" else (DATE if kind == "date" else
                                                (arg if kind == "calc" else None))
            if fmt:
                cell.number_format = fmt

    # 色：13項目は案件管理と同じ配色、判定は頭文字、服部判断は言葉で
    g_rng = f"B{ROW['g1-1']}:{last}{ROW['gエリア']}"
    for val, color in B.GRADE_COLORS:
        ws.conditional_formatting.add(g_rng, CellIsRule(
            operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", fgColor=color)))
    v_rng = f"B{ROW['ai']}:{last}{ROW['tomoi']}"
    for head, color in (("A", "B7E1CD"), ("B", "FFF2CC"), ("C", "F4C7C3")):
        ws.conditional_formatting.add(v_rng, FormulaRule(
            formula=[f'LEFT(B{ROW["ai"]},1)="{head}"'], fill=PatternFill("solid", fgColor=color)))
    j_rng = f"B{ROW['judge']}:{last}{ROW['judge']}"
    for val, color in zip(B.ESC_JUDGE, ("B7E1CD", "D9EAD3", "FFF2CC", "F4C7C3")):
        ws.conditional_formatting.add(j_rng, CellIsRule(
            operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", fgColor=color)))
    ws.conditional_formatting.add(f"B{ROW['m2']}:{last}{ROW['m2']}", CellIsRule(
        operator="greaterThan", formula=["7"], font=Font(color="C00000", bold=True)))

    ws[f"A{NOTE_ROW}"] = (
        "出所：案件管理・質問リスト（自動）。承継コストの概算＝仲介手数料（レーマン方式、最低20）＋DD"
        "（3億円以下10・10億円以下20・30億円以下30・それ超45）。"
        "TE目線の株式価値＝倍率×実態EBITDA−承継コスト＋実質NetCash（承継コストは希望価格で計算した額のまま）。"
        "この表の◯△×・判定・数字は案件管理の写しなので、直すときは案件管理の側を直す。")
    ws[f"A{NOTE_ROW}"].font = B.NOTE_F

    ws.column_dimensions["A"].width = 34
    for c in cols:
        ws.column_dimensions[c].width = 40
    ws.row_dimensions[R_HEAD].height = 30
    ws.freeze_panes = f"B{R_HEAD + 1}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "1F3864"
    return cols


# ── 現物への組み込み ─────────────────────────────────────────────────────
def load_spec(path):
    s = json.load(io.open(path, encoding="utf-8"))
    for d in s["deals"]:
        d["案件No"] = int(d["案件No"])
    return s


def to_date(s):
    y, m, d = (int(x) for x in s.split("/"))
    return datetime.datetime(y, m, d)


def fill(ws, col, rec, deal):
    ws[f"{col}{ROW['no']}"] = deal["案件No"]
    ws[f"{col}{ROW['date']}"] = to_date(rec["推薦日"])
    for key, field in (("point", "推薦ポイント"), ("issue", "主な論点")):
        ws[f"{col}{ROW[key]}"] = deal.get(field) or None
    if deal.get("承継コスト") not in (None, ""):
        ws[f"{col}{ROW['cost_in']}"] = float(deal["承継コスト"])


# ひな型（build_pipeline.py）に入れる記入例。案件管理の8行目（架空の No.1）を推薦した形
SAMPLE = {
    "推薦日": "2026/10/01",
    "deal": {"案件No": 1,
             "推薦ポイント": "（記入例）・重点①シニアの配食で、売上の8割が月額の定期契約\n"
                         "・3期連続黒字で増加基調、粗利率は3期とも30%以上\n・価格は②5.1倍で枠内",
             "主な論点": "（記入例）・得意先別の売上構成が未開示\n・社長の営業依存（顧問就任で緩和できるか）"},
    "推薦理由": "（記入例）重点業種に直球でストック性が高く、価格も枠内。トップ面談に進みたい",
    "服部判断": "条件付きで進める",
    "服部コメント": "（記入例）得意先の構成を確かめてから面談へ",
}


def fill_sample(ws, col):
    fill(ws, col, SAMPLE, SAMPLE["deal"])
    ws[f"{col}{ROW['why']}"] = SAMPLE["推薦理由"]
    ws[f"{col}{ROW['judge']}"] = SAMPLE["服部判断"]
    ws[f"{col}{ROW['comment']}"] = SAMPLE["服部コメント"]


def deal_numbers(wb):
    D = letters(wb[SH_DEAL], 7)
    ws = wb[SH_DEAL]
    out = set()
    for r in range(8, ws.max_row + 1):
        v = ws[f"{D['No']}{r}"].value
        try:
            out.add(int(float(v)))
        except (TypeError, ValueError):
            pass
    return out


def tab_columns(ws):
    """推薦案件タブの案件の列（見出しの行に式が入っている列）と、使用中の案件No。"""
    cols, used = [], {}
    for c in ws[R_HEAD][1:]:
        if is_formula(c.value):
            cols.append(c.column_letter)
            v = ws[f"{c.column_letter}{ROW['no']}"].value
            if v not in (None, ""):
                used[c.column_letter] = int(float(v))
    return cols, used


def check_rows(ws):
    """タブの行の並びがこの道具の想定どおりか（ラベルで確かめる）。"""
    for key, label, kind, _ in SPEC:
        if label is None or kind == "band":
            continue
        got = norm(ws[f"A{ROW[key]}"].value or "")
        if got != norm(label):
            raise SystemExit(f"推薦案件タブの{ROW[key]}行目が「{label}」でない（{got}）。"
                             "シート側で行を足した・消したときは tools/recommend.py の SPEC を合わせる")


def evaluate_tab(path, cols, maxrow):
    """保存した状態の推薦案件タブを評価して、案件ごとの要点とエラー数を返す。"""
    from xlformula import Book, Err, evaluate
    wb = load_workbook(path)
    b = Book(wb, maxrow=maxrow)
    ws = wb[SH]
    show = ["name", "ai", "tomoi", "rate", "m1", "m2", "cost", "lo", "hi", "openq", "mustq"]
    errs, rows = 0, []
    for c in cols:
        got = {}
        for key, _, kind, _ in SPEC:
            v = ws[f"{c}{ROW[key]}"].value
            if not is_formula(v):
                continue
            try:
                got[key] = evaluate(b, SH, v.text if isinstance(v, ArrayFormula) else v)
            except Err as e:
                errs += 1
                got[key] = f"#ERR {e.kind}"
        if ws[f"{c}{ROW['no']}"].value not in (None, ""):
            rows.append((c, ws[f"{c}{ROW['no']}"].value, {k: got.get(k) for k in show}))
    return errs, rows


def print_rows(rows):
    for c, no, g in rows:
        f = lambda v, d=1: (f"{v:.{d}f}" if isinstance(v, float) else v)  # noqa: E731
        print(f"   {c}列 No.{int(no)} {g['name']} | {g['ai']} | 友井 {g['tomoi'] or '－'} | "
              f"達成率 {f(g['rate'], 3)} | ① {f(g['m1'])} ② {f(g['m2'])} | 承継コスト {f(g['cost'])} | "
              f"目線 {f(g['lo'])}〜{f(g['hi'])} | 未解決 {g['openq']}（必須 {g['mustq']}）")


def cmd_tab(a):
    rec = load_spec(a.deals)
    wb = load_workbook(a.layout)
    if SH in wb.sheetnames:
        raise SystemExit(f"{SH}タブはもうある。案件を足すときは blocks を使う")
    known = deal_numbers(wb)
    D, Q = letters(wb[SH_DEAL], 7), letters(wb[SH_Q], 3)
    ws = wb.create_sheet(SH, wb.sheetnames.index(SH_DEAL) + 1)
    cols = build_recommend(ws, D, Q)
    for col, deal in zip(cols, rec["deals"]):
        if deal["案件No"] not in known:
            raise SystemExit(f"案件No {deal['案件No']} が案件管理に無い")
        fill(ws, col, rec, deal)
    wb.calculation.fullCalcOnLoad = True
    wb.save(a.out)
    print(f"{a.out} に {SH}タブを追加（{len(rec['deals'])}件、{cols[0]}〜{cols[len(rec['deals']) - 1]}列）")
    if a.verify:
        errs, rows = evaluate_tab(a.out, cols, maxrow=a.maxrow)
        print_rows(rows)
        print(f"── {SH}タブの数式エラー: {errs} 件")
        if errs:
            raise SystemExit("数式エラーあり")


def one_line(s):
    """貼り付けブロックの1セルに入れる文。改行はセルの区切りと取り違えるので全角空白にする。"""
    return re.sub(r"\s*\n\s*", "　", (s or "").strip()).replace("\t", " ")


def cmd_blocks(a):
    rec = load_spec(a.deals)
    wb = load_workbook(a.layout)
    if SH not in wb.sheetnames:
        raise SystemExit(f"現物に{SH}タブが無い。先に tab でタブを作る")
    ws = wb[SH]
    check_rows(ws)
    known = deal_numbers(wb)
    cols, used = tab_columns(ws)
    free = [c for c in cols if c not in used]
    deals = rec["deals"]
    if len(free) < len(deals):
        raise SystemExit(f"{SH}タブの空き列が{len(free)}列しかない。最後の列（{cols[-1]}）を右へ"
                         "コピーして列を増やし、ダウンロードし直してから作り直す")
    out = io.StringIO()
    out.write(f"【{SH}タブに案件を足す】{len(deals)}件\n"
              "指定のセルを1つクリックして貼るだけ。式の入った行には触れない。\n"
              "推薦理由・服部判断・服部コメントは黄色のセルに直接書いてください。\n\n")
    n = 0
    for col, d in zip(free, deals):
        no = d["案件No"]
        if no not in known:
            raise SystemExit(f"案件No {no} が案件管理に無い")
        if no in used.values():
            print(f"   注意：案件No {no} はもう{SH}タブにある（{[c for c, v in used.items() if v == no]}列）")
        n += 1
        out.write(f"━━━ ブロック{n}：「{SH}」の {col}{ROW['no']} をクリックして貼る（1列 × 4行）\n"
                  "  行順：案件No / 推薦日 / 推薦ポイント / 主な論点・リスク\n")
        out.write("\n".join([str(no), rec["推薦日"], one_line(d.get("推薦ポイント")),
                             one_line(d.get("主な論点"))]) + "\n\n")
        if d.get("承継コスト") not in (None, ""):
            n += 1
            out.write(f"━━━ ブロック{n}：「{SH}」の {col}{ROW['cost_in']} をクリックして貼る（1セル）\n"
                      f"  承継コスト（仲介の試算など）\n{d['承継コスト']}\n\n")
    io.open(a.out, "w", encoding="utf-8").write(out.getvalue())
    print(f"{a.out} を作成：{len(deals)}件 / {free[0]}列から / ブロック{n}個")
    if a.verify:
        tmp = a.save or "_recommend_verify.xlsx"
        shutil.copy(a.layout, tmp)
        wb2 = load_workbook(tmp)
        for col, d in zip(free, deals):
            dd_ = dict(d, 推薦ポイント=one_line(d.get("推薦ポイント")),
                       主な論点=one_line(d.get("主な論点")))
            fill(wb2[SH], col, rec, dd_)
        wb2.save(tmp)
        errs, rows = evaluate_tab(tmp, cols, maxrow=a.maxrow)
        print_rows([r for r in rows if r[0] in free[:len(deals)]])
        print(f"── {SH}タブの数式エラー: {errs} 件")
        if not a.save:
            os.remove(tmp)
        if errs:
            raise SystemExit("数式エラーあり")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("tab", "blocks"):
        p = sub.add_parser(name)
        p.add_argument("--layout", required=True, help="ドライブからダウンロードした案件管理表（xlsx）")
        p.add_argument("--deals", required=True, help="推薦日と案件を書いたJSON")
        p.add_argument("--out", required=True)
        p.add_argument("--verify", action="store_true")
        p.add_argument("--maxrow", type=int, default=600,
                       help="列まるごとの参照を何行目まで評価するか（質問リストの行数より大きく）")
        if name == "blocks":
            p.add_argument("--save", metavar="XLSX", help="検証に使った貼り付け後の状態を残す")
    a = ap.parse_args()
    (cmd_tab if a.cmd == "tab" else cmd_blocks)(a)


if __name__ == "__main__":
    main()
