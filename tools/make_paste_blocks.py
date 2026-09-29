#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""案件管理表（Googleスプレッドシート）へ直接貼るブロックを作る。

Apps Script は使わない運用。列の並びは**ドライブの現物をダウンロードした
xlsx から読む**ので、シート側で列を増やしてもブロックの宛先がずれない。

  python3 tools/make_paste_blocks.py --layout live.xlsx --paste paste.json \
          --out blocks.txt [--verify]

paste.json
  {
    "deal": {"企業名": "...", "譲渡価格": 450, "★1-1": "◯", ...},
    "questions": [["分類", "質問", "何を確かめたいか", "関連項目", "優先度"], ...]
  }
  ・deal のキーは案件管理の見出しそのまま（改行は詰める。★も付けたまま）
  ・No は書かない（次の番号を自動で振る）
  ・questions の案件No と Q# も自動で入る

--verify を付けると、貼った後の状態を作って全数式を評価し、
総合判定・達成率・倍率まで出して確かめる。
"""
import argparse
import datetime
import io
import json
import os
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

SH_DEAL, SH_Q = "案件管理", "質問リスト"
DEAL_HEAD, DEAL_FIRST, DEAL_LAST = 4, 5, 104
Q_HEAD, Q_FIRST, Q_LAST = 3, 4, 303
DATE_COLS = {"流入日", "期限", "意向表明期限", "聞いた日", "回答日"}
Q_FIELDS = ["分類", "質問（このまま読める文）", "何を確かめたいか", "関連項目", "優先度"]


def is_formula(v):
    return isinstance(v, ArrayFormula) or (isinstance(v, str) and v.startswith("="))


def formula_text(v):
    return v.text if isinstance(v, ArrayFormula) else v


def head_map(ws, row):
    out = {}
    for c in ws[row]:
        if c.value is None:
            continue
        out.setdefault(str(c.value).replace("\n", "").strip(), c.column)
    return out


def input_runs(ws, head_row, probe_row):
    """数式が入っていない列を、連続する塊に分ける。塊ごとに1回貼る。"""
    cols = sorted(head_map(ws, head_row).values())
    inp = [c for c in cols if not is_formula(ws.cell(probe_row, c).value)]
    runs = []
    for c in inp:
        if runs and c == runs[-1][-1] + 1:
            runs[-1].append(c)
        else:
            runs.append([c])
    return runs


def next_deal_row(ws, no_col):
    row = DEAL_FIRST
    mx = 0
    for r in range(DEAL_FIRST, DEAL_LAST + 1):
        v = ws.cell(r, no_col).value
        if v not in ("", None) and not is_formula(v):
            row = r + 1
            try:
                mx = max(mx, float(v))
            except (TypeError, ValueError):
                pass
    return row, int(mx) + 1


def next_q_row(ws, no_col, seq_col, deal_no):
    row, seq = Q_FIRST, 0
    for r in range(Q_FIRST, Q_LAST + 1):
        v = ws.cell(r, no_col).value
        if v in ("", None) or is_formula(v):
            continue
        row = r + 1
        try:
            if float(v) == float(deal_no):
                seq = max(seq, float(ws.cell(r, seq_col).value or 0))
        except (TypeError, ValueError):
            pass
    return row, int(seq)


def coerce(head, raw):
    if raw in ("", None):
        return ""
    if head in DATE_COLS and isinstance(raw, str) and "/" in raw:
        return raw
    return raw


def build(layout, paste):
    wb = load_workbook(layout)
    wd, wq = wb[SH_DEAL], wb[SH_Q]
    dh, qh = head_map(wd, DEAL_HEAD), head_map(wq, Q_HEAD)
    inv_d = {v: k for k, v in dh.items()}
    inv_q = {v: k for k, v in qh.items()}

    unknown = [k for k in paste["deal"] if k not in dh]
    if unknown:
        raise SystemExit(f"案件管理に無い見出しが deal に含まれています: {unknown}")

    drow, dno = next_deal_row(wd, dh["No"])
    if drow > DEAL_LAST:
        raise SystemExit("案件管理の行が埋まっています。行を足してください。")
    qrow, qseq = next_q_row(wq, qh["案件No"], qh["Q#"], dno)
    qs = paste.get("questions", [])
    if qrow + len(qs) - 1 > Q_LAST:
        raise SystemExit("質問リストの行が埋まっています。行を足してください。")

    deal = dict(paste["deal"])
    deal["No"] = dno
    qrows = []
    for i, q in enumerate(qs):
        m = dict(zip(Q_FIELDS, q))
        m["案件No"] = dno
        m["Q#"] = qseq + i + 1
        m["状態"] = "未質問"
        qrows.append(m)

    out = io.StringIO()
    out.write("【案件管理表への貼り付け】Apps Scriptは使いません。\n")
    out.write(f"案件No {dno}「{deal.get('企業名','')}」／質問 {len(qs)}件\n")
    out.write("数式が入っている列は各ブロックから外してあるので、"
              "指定セルを1つクリックして貼れば数式は壊れません。\n")
    out.write("貼る順番は問いません。\n\n")

    n = 0
    for run in input_runs(wd, DEAL_HEAD, drow):
        n += 1
        addr = f"{get_column_letter(run[0])}{drow}"
        cols = [inv_d[c] for c in run]
        out.write(f"━━━ ブロック{n}：「{SH_DEAL}」の {addr} をクリックして貼る"
                  f"（{len(run)}列 × 1行）\n")
        out.write("  列順：" + " / ".join(cols) + "\n")
        out.write("\t".join(str(coerce(h, deal.get(h, ""))) for h in cols) + "\n\n")
    for run in input_runs(wq, Q_HEAD, qrow):
        if not qrows:
            break
        n += 1
        addr = f"{get_column_letter(run[0])}{qrow}"
        cols = [inv_q[c] for c in run]
        out.write(f"━━━ ブロック{n}：「{SH_Q}」の {addr} をクリックして貼る"
                  f"（{len(run)}列 × {len(qrows)}行）\n")
        out.write("  列順：" + " / ".join(cols) + "\n")
        for m in qrows:
            out.write("\t".join(str(coerce(h, m.get(h, ""))) for h in cols) + "\n")
        out.write("\n")
    return out.getvalue(), dict(deal_row=drow, deal_no=dno, q_row=qrow,
                                q_count=len(qrows), blocks=n)


def verify(layout, paste, info):
    """ブロックを実際に貼った状態を作り、全数式を評価して結果を出す。"""
    from xlformula import Book, evaluate, Err
    tmp = "_paste_verify.xlsx"
    shutil.copy(layout, tmp)
    wb = load_workbook(tmp)
    wd, wq = wb[SH_DEAL], wb[SH_Q]
    dh, qh = head_map(wd, DEAL_HEAD), head_map(wq, Q_HEAD)
    drow, dno, qrow = info["deal_row"], info["deal_no"], info["q_row"]

    def put(ws, head, row, h, v):
        if v in ("", None):
            return
        col = head[h]
        if h in DATE_COLS and isinstance(v, str) and "/" in v:
            y, m, d = (int(x) for x in v.split("/"))
            ws.cell(row, col).value = datetime.date(y, m, d)
        else:
            ws.cell(row, col).value = v

    put(wd, dh, drow, "No", dno)
    for h, v in paste["deal"].items():
        put(wd, dh, drow, h, v)
    for i, q in enumerate(paste.get("questions", [])):
        m = dict(zip(Q_FIELDS, q))
        m["案件No"], m["Q#"], m["状態"] = dno, i + 1, "未質問"
        for h, v in m.items():
            put(wq, qh, qrow + i, h, v)
    wb.save(tmp)

    wb2 = load_workbook(tmp)
    b = Book(wb2, maxrow=Q_LAST + 20)
    ws, D = wb2[SH_DEAL], head_map(wb2[SH_DEAL], DEAL_HEAD)
    print(f"── 貼り付け後の {SH_DEAL} {drow}行目 ──")
    for h in ("No", "企業名", "EV/EBITDAマルチプル", "参考：簿価ベース倍率",
              "売上レンジ", "総合判定", "達成率", "未解決質問"):
        if h not in D:
            continue
        v = ws.cell(drow, D[h]).value
        if is_formula(v):
            v = evaluate(b, SH_DEAL, formula_text(v))
        print(f"   {h:<22} {round(v, 3) if isinstance(v, float) else v}")
    errs = 0
    for name in wb2.sheetnames:
        for row in wb2[name].iter_rows():
            for c in row:
                if is_formula(c.value):
                    try:
                        evaluate(b, name, formula_text(c.value))
                    except Err:
                        errs += 1
    print(f"── 全数式のエラー: {errs} 件")
    os.remove(tmp)
    return errs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--layout", required=True,
                    help="ドライブからダウンロードした案件管理表（xlsx）")
    ap.add_argument("--paste", required=True, help="deal と questions を書いたJSON")
    ap.add_argument("--out", required=True)
    ap.add_argument("--verify", action="store_true")
    a = ap.parse_args()
    paste = json.load(io.open(a.paste, encoding="utf-8"))
    text, info = build(a.layout, paste)
    io.open(a.out, "w", encoding="utf-8").write(text)
    print(f"{a.out} を作成：案件No {info['deal_no']} / "
          f"{SH_DEAL} {info['deal_row']}行目 / {SH_Q} {info['q_row']}行目から"
          f"{info['q_count']}件 / ブロック{info['blocks']}個")
    if a.verify:
        if verify(a.layout, paste, info):
            raise SystemExit("数式エラーあり")


if __name__ == "__main__":
    main()
