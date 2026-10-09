#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""案件管理表に既にある案件を、届いた資料で更新するブロックを作る。

新しい案件は make_paste_blocks.py。こちらは「No.6 の概要書が届いた」
「決算書をもらった」のように、行が既にある案件の上書き・回答の記入・
質問の追加をまとめて出す。宛先は layout（ドライブの現物）から引くので手で数えない。

  python3 tools/update_deal_blocks.py --layout live.xlsx --spec spec.json \
          --out blocks.txt [--verify [--save XLSX]]

spec.json
  {
    "案件No": 6,
    "deal": {"企業名": "...", "実態EBITDA": 158.4, "★1-3": "◯", ...},   ← 変える列だけ
    "answers": [                                                        ← 案件No × Q# で行を引く
      {"Q": 3, "回答": "...", "回答日": "2026/10/09", "回答で動く評価": "5", "状態": "解決"},
      {"Q": 9, "回答": "...", "状態": "追加確認が必要", "追記": true}     ← 前の回答の後ろに足す
    ],
    "出所": "レコフ 企業概要書v5（10/9受領）",                            ← 回答の末尾に付ける
    "questions": [["分類", "質問", "何を確かめたいか", "関連項目", "優先度"], ...]  ← 末尾に追加
  }

・deal に数式の列や「友井判定」を書こうとしたら止まる（貼ると壊す／上書きする）
・塊の中で変えない列は今の値をそのまま入れるので、1回の貼り付けで済む
"""
import argparse
import datetime
import io
import json
import os
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import make_paste_blocks as mpb  # noqa: E402

ANS_COLS = ["聞いた日", "回答", "回答日", "回答で動く評価", "状態"]
STATES = ["未質問", "回答待ち", "回答済", "追加確認が必要", "解決"]


def cell_text(h, v):
    """今の値をブロックに書き戻すときの文字列。日付は 2026/10/09 の形にする。"""
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y/%m/%d")
    if isinstance(v, float) and v.is_integer() and h not in ("実態EBITDA", "譲渡価格"):
        return str(int(v))
    return str(v)


def deal_row(wd, dh, no):
    for r in range(mpb.DEAL_FIRST, mpb.DEAL_LAST + 1):
        v = wd.cell(r, dh["No"]).value
        try:
            if v not in (None, "") and float(v) == float(no):
                return r
        except (TypeError, ValueError):
            pass
    raise SystemExit(f"案件管理に No.{no} の行が無い")


def q_rows(wq, qh, no):
    out = {}
    for r in range(mpb.Q_FIRST, mpb.Q_LAST + 1):
        v = wq.cell(r, qh["案件No"]).value
        try:
            if v not in (None, "") and float(v) == float(no):
                out[int(wq.cell(r, qh["Q#"]).value)] = r
        except (TypeError, ValueError):
            pass
    return out


def contiguous(rows):
    runs = []
    for r in sorted(rows):
        if runs and r == runs[-1][-1] + 1:
            runs[-1].append(r)
        else:
            runs.append([r])
    return runs


def plan(layout, spec):
    """書き込む中身を (シート, 行, 見出し) → 値 で作る。ブロックと検算の両方がこれを使う。"""
    wb = load_workbook(layout)
    mpb.set_limits(wb)
    wd, wq = wb[mpb.SH_DEAL], wb[mpb.SH_Q]
    dh, qh = mpb.head_map(wd, mpb.DEAL_HEAD), mpb.head_map(wq, mpb.Q_HEAD)
    no = spec["案件No"]
    drow = deal_row(wd, dh, no)

    writes = {}
    for h, v in (spec.get("deal") or {}).items():
        if h not in dh:
            raise SystemExit(f"案件管理に無い見出し: {h}")
        if h in mpb.SKIP_COLS or h == "No":
            raise SystemExit(f"{h} は書かない列")
        if mpb.is_formula(wd.cell(drow, dh[h]).value):
            raise SystemExit(f"{h} は数式の列なので書けない")
        writes[(mpb.SH_DEAL, drow, h)] = v

    rows = q_rows(wq, qh, no)
    src = spec.get("出所", "")
    for a in spec.get("answers") or []:
        if a["Q"] not in rows:
            raise SystemExit(f"No.{no} の Q{a['Q']} が質問リストに無い")
        if a["状態"] not in STATES:
            raise SystemExit(f"状態が選択肢外: {a['状態']}")
        r = rows[a["Q"]]
        cur = wq.cell(r, qh["回答"]).value
        body = a["回答"] + (f"【出所】{src}" if src and src not in a["回答"] else "")
        if a.get("追記"):
            d = a.get("回答日", "")
            tag = f"【{int(d.split('/')[1])}/{int(d.split('/')[2])}追記】" if d else "【追記】"
            body = f"{cur}／{tag}{body}" if cur else f"{tag}{body}"
        elif cur and cur != body:
            raise SystemExit(f"Q{a['Q']}（行{r}）には既に回答がある。書き足すなら \"追記\": true")
        writes[(mpb.SH_Q, r, "回答")] = body
        writes[(mpb.SH_Q, r, "状態")] = a["状態"]
        for h in ("聞いた日", "回答日", "回答で動く評価"):
            if a.get(h):
                writes[(mpb.SH_Q, r, h)] = a[h]

    new = spec.get("questions") or []
    qrow = None
    if new:
        qrow, _ = mpb.next_q_row(wq, qh["案件No"], qh["Q#"], no)
        if qrow + len(new) - 1 > mpb.Q_LAST:
            raise SystemExit(f"質問リストの行が足りません（数式は{mpb.Q_LAST}行目まで）")
        qmax = max(rows) if rows else 0
        for i, q in enumerate(new):
            m = dict(zip(mpb.Q_FIELDS, q))
            m.update({"案件No": no, "Q#": qmax + 1 + i, "状態": "未質問"})
            for h, v in m.items():
                writes[(mpb.SH_Q, qrow + i, h)] = v
        reach = mpb.openq_reach(wd, dh, mpb.DEAL_FIRST)
        if reach is not None and qrow + len(new) - 1 > reach:
            print(f"【注意】未解決質問の数式は質問リスト{reach}行目までしか数えていません。"
                  "列全体（$A:$A・$O:$O）に直してから貼ってください。", file=sys.stderr)
    return wb, dict(no=no, deal_row=drow, q_row=qrow, new=len(new),
                    answered=len(spec.get("answers") or [])), writes


def blocks(wb, info, writes):
    wd, wq = wb[mpb.SH_DEAL], wb[mpb.SH_Q]
    dh, qh = mpb.head_map(wd, mpb.DEAL_HEAD), mpb.head_map(wq, mpb.Q_HEAD)
    inv_d = {v: k for k, v in dh.items()}
    inv_q = {v: k for k, v in qh.items()}
    out, n = io.StringIO(), 0

    def emit(sheet, ws, inv, run, rows, label):
        nonlocal n
        n += 1
        cols = [inv[c] for c in run]
        addr = f"{get_column_letter(run[0])}{rows[0]}"
        out.write(f"━━━ ブロック{n}：「{sheet}」の {addr} をクリックして貼る"
                  f"（{len(run)}列 × {len(rows)}行・{label}）\n")
        out.write("  列順：" + " / ".join(cols) + "\n")
        for r in rows:
            vals = []
            for c, h in zip(run, cols):
                v = writes.get((sheet, r, h), ws.cell(r, c).value)
                vals.append(cell_text(h, v))
            out.write("\t".join(vals) + "\n")
        out.write("\n")

    # 案件管理：変える列を含む塊だけ。No 列は外す
    drow = info["deal_row"]
    no_col = dh["No"]
    for run in mpb.input_runs(wd, mpb.DEAL_HEAD, drow):
        run = [c for c in run if c != no_col]
        if run and any((mpb.SH_DEAL, drow, inv_d[c]) in writes for c in run):
            emit(mpb.SH_DEAL, wd, inv_d, run, [drow], "上書き")

    # 質問リスト：回答の5列（聞いた日〜状態）を、連続する行ごとに
    ans_cols = [qh[h] for h in ANS_COLS]
    if ans_cols != list(range(ans_cols[0], ans_cols[0] + len(ans_cols))):
        raise SystemExit("回答の5列が隣り合っていない。シートの列順を確認")
    new_rows = set()
    if info["new"]:
        new_rows = set(range(info["q_row"], info["q_row"] + info["new"]))
    ans_rows = {r for (s, r, h) in writes if s == mpb.SH_Q and h in ANS_COLS} - new_rows
    for rows in contiguous(ans_rows):
        emit(mpb.SH_Q, wq, inv_q, ans_cols, rows, "回答の上書き")

    # 質問リスト：追加の質問（数式の列は外す）
    if new_rows:
        rows = sorted(new_rows)
        for run in mpb.input_runs(wq, mpb.Q_HEAD, rows[0]):
            emit(mpb.SH_Q, wq, inv_q, run, rows, "新規")
    return out.getvalue(), n


def verify(layout, wb_plan, info, writes, save=None):
    from xlformula import Book, evaluate, Err
    tmp = save or "_update_verify.xlsx"
    shutil.copy(layout, tmp)
    wb = load_workbook(tmp)
    heads = {s: mpb.head_map(wb[s], mpb.DEAL_HEAD if s == mpb.SH_DEAL else mpb.Q_HEAD)
             for s in (mpb.SH_DEAL, mpb.SH_Q)}
    for (s, r, h), v in writes.items():
        if isinstance(v, str) and h in mpb.DATE_COLS and "/" in v:
            y, m, d = (int(x) for x in v.split("/"))
            v = datetime.date(y, m, d)
        wb[s].cell(r, heads[s][h]).value = v
    wb.save(tmp)
    wb2 = load_workbook(tmp)
    b = Book(wb2, maxrow=mpb.Q_LAST + 20)
    ws, D = wb2[mpb.SH_DEAL], heads[mpb.SH_DEAL]
    r = info["deal_row"]
    print(f"── 更新後の {mpb.SH_DEAL} {r}行目 ──")
    for h in ("No", "企業名", "EV/EBITDAマルチプル", "参考：簿価ベース倍率", "売上レンジ",
              "AI総合判定", "達成率", "未解決質問"):
        if h not in D:
            print(f"   【見出しが見つからない】{h}")
            continue
        v = ws.cell(r, D[h]).value
        if mpb.is_formula(v):
            v = evaluate(b, mpb.SH_DEAL, mpb.formula_text(v))
        print(f"   {h:<22} {round(v, 3) if isinstance(v, float) else v}")
    errs = 0
    for name in wb2.sheetnames:
        for row in wb2[name].iter_rows():
            for c in row:
                if mpb.is_formula(c.value):
                    try:
                        evaluate(b, name, mpb.formula_text(c.value))
                    except Err:
                        errs += 1
    print(f"── 全数式のエラー: {errs} 件")
    if save:
        print(f"── 更新済みの状態を {save} に保存した")
    else:
        os.remove(tmp)
    return errs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--layout", required=True)
    ap.add_argument("--spec", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--verify", action="store_true")
    ap.add_argument("--save", metavar="XLSX")
    a = ap.parse_args()
    if a.save and not a.verify:
        raise SystemExit("--save は --verify と一緒に指定する")
    spec = json.load(io.open(a.spec, encoding="utf-8"))
    wb, info, writes = plan(a.layout, spec)
    text, n = blocks(wb, info, writes)
    head = (f"【案件管理表の更新】案件No.{info['no']}（案件管理 {info['deal_row']}行目）\n"
            f"回答の記入 {info['answered']}件／質問の追加 {info['new']}件\n"
            "数式の列は外してあるので、指定セルを1つクリックして貼れば数式は壊れません。\n"
            "塊の中で変えない列には今の値が入っています。\n\n")
    io.open(a.out, "w", encoding="utf-8").write(head + text)
    print(f"{a.out} を作成：ブロック{n}個")
    if a.verify and verify(a.layout, wb, info, writes, save=a.save):
        raise SystemExit("数式エラーあり")


if __name__ == "__main__":
    main()
